"""
Sistema de Gestao de Estoque com 
Frontend: CustomTkinter
Backend: SQLite (Gestao, Usuarios, Vendas)
Updates: Alerta por e-mail e relatorio Excel
"""

import sqlite3
import hashlib
import smtplib
from email.mime.text import MIMEText
from datetime import datetime
import os

import openpyxl
from openpyxl.chart import BarChart, Reference

import customtkinter as ctk
from tkinter import messagebox, ttk

ESTOQUE_ALERTA_EMAIL = 10

EMAIL_REMETENTE = "" #EMAIL
EMAIL_SENHA = "" #SENHA DE APP
SMTP_SERVIDOR = "smtp.gmail.com"
SMTP_PORTA = 587

#Alerta Email
def enviar_alerta_estoque(destinatario, produto, quantidade):
    assunto = f"Alerta de estoque baixo: {produto}"
    corpo = (
        f"O produto '{produto}' esta com estoque baixo.\n"
        f"Quantidade atual: {quantidade} unidades.\n\n"
        f"Sistema de Gestao de Estoque"
    )

    msg = MIMEText(corpo)
    msg["Subject"] = assunto
    msg["From"] = EMAIL_REMETENTE
    msg["To"] = destinatario

    try:
        servidor = smtplib.SMTP(SMTP_SERVIDOR, SMTP_PORTA)
        servidor.starttls()
        servidor.login(EMAIL_REMETENTE, EMAIL_SENHA)
        servidor.sendmail(EMAIL_REMETENTE, destinatario, msg.as_string())
        servidor.quit()
        print(f"E-mail de alerta enviado para {destinatario}")
    except Exception as erro:
        print(f"Erro ao enviar e-mail: {erro}")


def verificar_e_alertar(gestao, produto, email_destino):
    quantidade_atual = gestao.consultar_estoque(produto)
    if 0 < quantidade_atual <= ESTOQUE_ALERTA_EMAIL:
        enviar_alerta_estoque(email_destino, produto, quantidade_atual)

class Gestao:
    def __init__(self, banco):
        self.conn = sqlite3.connect(banco)
        self.criar_tabela_estoque()
        self.criar_tabela_fornecedores()

    def criar_tabela_estoque(self):
        cursor = self.conn.cursor()
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS estoque (
        id INTEGER PRIMARY KEY,
        produto TEXT UNIQUE,
        quantidade INTEGER,
        preco REAL
        )
        ''')
        self.conn.commit()

    def criar_tabela_fornecedores(self):
        cursor = self.conn.cursor()
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS fornecedores (
        id INTEGER PRIMARY KEY,
        nome TEXT,
        celular TEXT,
        produto TEXT
        )
        ''')
        self.conn.commit()

    def adicionar_produto(self, produto, quantidade, preco=None):
        cursor = self.conn.cursor()
        cursor.execute("SELECT quantidade, preco FROM estoque WHERE produto=?", (produto,))
        resultado = cursor.fetchone()
        if resultado:
            nova_quantidade = resultado[0] + quantidade
            preco_final = preco if preco is not None else resultado[1]
            cursor.execute("UPDATE estoque SET quantidade=?, preco=? WHERE produto=?",
                            (nova_quantidade, preco_final, produto))
        else:
            if preco is None:
                preco = 0
            nova_quantidade = quantidade
            cursor.execute("INSERT INTO estoque (produto, quantidade, preco) VALUES (?,?,?)",
                            (produto, quantidade, preco))
        self.conn.commit()
        return nova_quantidade

    def consultar_produto_por_id(self, produto_id):
        cursor = self.conn.cursor()
        cursor.execute("SELECT id, produto, quantidade, preco FROM estoque WHERE id=?", (produto_id,))
        return cursor.fetchone()

    def adicionar_produto_por_id(self, produto_id, quantidade, preco=None):
        produto_existente = self.consultar_produto_por_id(produto_id)
        if not produto_existente:
            return False, None

        _, produto, quantidade_atual, preco_atual = produto_existente
        nova_quantidade = quantidade_atual + quantidade
        preco_final = preco if preco is not None else preco_atual

        cursor = self.conn.cursor()
        cursor.execute("UPDATE estoque SET quantidade=?, preco=? WHERE id=?",
                        (nova_quantidade, preco_final, produto_id))
        self.conn.commit()
        return True, produto

    def executar_venda(self, produto, quantidade):
        cursor = self.conn.cursor()
        cursor.execute("SELECT quantidade FROM estoque WHERE produto=?", (produto,))
        resultado = cursor.fetchone()
        if resultado:
            estoque_atual = resultado[0]
            if estoque_atual >= quantidade:
                nova_quantidade = estoque_atual - quantidade
                cursor.execute("UPDATE estoque SET quantidade=? WHERE produto=?",
                                (nova_quantidade, produto))
                self.conn.commit()
                return True
            return False
        return False

    def remover_produto(self, produto):
        cursor = self.conn.cursor()
        cursor.execute("SELECT produto FROM estoque WHERE produto=?", (produto,))
        if not cursor.fetchone():
            return False
        cursor.execute("DELETE FROM estoque WHERE produto=?", (produto,))
        self.conn.commit()
        return True

    def consultar_estoque(self, produto):
        cursor = self.conn.cursor()
        cursor.execute("SELECT quantidade FROM estoque WHERE produto=?", (produto,))
        resultado = cursor.fetchone()
        return resultado[0] if resultado else 0

    def consultar_preco(self, produto):
        cursor = self.conn.cursor()
        cursor.execute("SELECT preco FROM estoque WHERE produto=?", (produto,))
        resultado = cursor.fetchone()
        return resultado[0] if resultado else None

    def lista_estoque(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT produto, quantidade, preco FROM estoque WHERE quantidade > 0")
        return cursor.fetchall()

    def lista_estoque_completa(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT id, produto, quantidade, preco FROM estoque")
        return cursor.fetchall()

    def valor_total_estoque(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT SUM(quantidade * preco) FROM estoque")
        resultado = cursor.fetchone()
        return resultado[0] if resultado[0] is not None else 0

    def gerar_grafico_estoque(self, caminho_saida="relatorio_estoque.xlsx"):
        produtos = self.lista_estoque()
        if not produtos:
            return False

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Estoque"
        ws.append(["Produto", "Quantidade", "Preco"])
        for produto, quantidade, preco in produtos:
            ws.append([produto, quantidade, preco])

        ultima_linha = ws.max_row

        grafico = BarChart()
        grafico.type = "col"
        grafico.title = "Quantidade em Estoque por Produto"
        grafico.x_axis.title = "Produto"
        grafico.y_axis.title = "Quantidade"
        grafico.width = 20
        grafico.height = 12

        dados_ref = Reference(ws, min_col=2, min_row=1, max_row=ultima_linha)
        categorias_ref = Reference(ws, min_col=1, min_row=2, max_row=ultima_linha)
        grafico.add_data(dados_ref, titles_from_data=True)
        grafico.set_categories(categorias_ref)
        ws.add_chart(grafico, "E2")

        wb.save(caminho_saida)
        return True

    def adicionar_fornecedor(self, nome, celular, produto):
        cursor = self.conn.cursor()
        cursor.execute(
            "INSERT INTO fornecedores (nome, celular, produto) VALUES (?,?,?)",
            (nome, celular, produto)
        )
        self.conn.commit()

    def listar_fornecedores(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT id, nome, celular, produto FROM fornecedores")
        return cursor.fetchall()

    def remover_fornecedor(self, fornecedor_id):
        cursor = self.conn.cursor()
        cursor.execute("SELECT nome FROM fornecedores WHERE id=?", (fornecedor_id,))
        resultado = cursor.fetchone()
        if not resultado:
            return None
        cursor.execute("DELETE FROM fornecedores WHERE id=?", (fornecedor_id,))
        self.conn.commit()
        return resultado[0]


class Usuarios:
    def __init__(self, banco):
        self.conn = sqlite3.connect(banco)
        self.criar_tabela_usuarios()

    def criar_tabela_usuarios(self):
        cursor = self.conn.cursor()
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY,
        username TEXT UNIQUE,
        senha_hash TEXT,
        email TEXT
        )
        ''')
        self.conn.commit()

    def _gerar_hash(self, senha):
        return hashlib.sha256(senha.encode()).hexdigest()

    def cadastrar_usuario(self, username, senha, email):
        cursor = self.conn.cursor()
        cursor.execute("SELECT username FROM usuarios WHERE username=?", (username,))
        if cursor.fetchone():
            return False
        senha_hash = self._gerar_hash(senha)
        cursor.execute(
            "INSERT INTO usuarios (username, senha_hash, email) VALUES (?,?,?)",
            (username, senha_hash, email)
        )
        self.conn.commit()
        return True

    def login(self, username, senha):
        cursor = self.conn.cursor()
        senha_hash = self._gerar_hash(senha)
        cursor.execute(
            "SELECT id, username, email FROM usuarios WHERE username=? AND senha_hash=?",
            (username, senha_hash)
        )
        resultado = cursor.fetchone()
        if resultado:
            id_usuario, username, email = resultado
            return {"id": id_usuario, "username": username, "email": email}
        return None


class Vendas:
    def __init__(self, banco, gestao):
        self.conn = sqlite3.connect(banco)
        self.gestao = gestao
        self.criar_tabela_vendas()

    def criar_tabela_vendas(self):
        cursor = self.conn.cursor()
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS vendas (
        id INTEGER PRIMARY KEY,
        produto TEXT,
        quantidade INTEGER,
        preco_unitario REAL,
        valor_total REAL,
        data TEXT
        )
        ''')
        self.conn.commit()

    def registrar_venda(self, produto, quantidade):
        preco = self.gestao.consultar_preco(produto)
        if preco is None:
            return False, "Produto nao encontrado no estoque."

        estoque_atual = self.gestao.consultar_estoque(produto)
        if estoque_atual < quantidade:
            return False, f"Estoque insuficiente (disponivel: {estoque_atual})."

        sucesso = self.gestao.executar_venda(produto, quantidade)
        if not sucesso:
            return False, "Falha ao executar a venda."

        valor_total = preco * quantidade
        data = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        cursor = self.conn.cursor()
        cursor.execute(
            "INSERT INTO vendas (produto, quantidade, preco_unitario, valor_total, data) VALUES (?,?,?,?,?)",
            (produto, quantidade, preco, valor_total, data)
        )
        self.conn.commit()
        return True, valor_total

    def remover_venda(self, venda_id):
        cursor = self.conn.cursor()
        cursor.execute("SELECT produto, quantidade FROM vendas WHERE id=?", (venda_id,))
        resultado = cursor.fetchone()
        if not resultado:
            return False
        produto, quantidade = resultado
        self.gestao.adicionar_produto(produto, quantidade)
        cursor.execute("DELETE FROM vendas WHERE id=?", (venda_id,))
        self.conn.commit()
        return True

    def listar_vendas(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT id, produto, quantidade, preco_unitario, valor_total, data FROM vendas")
        return cursor.fetchall()


#Cores:
COR_FUNDO = "#0d1117"
COR_SIDEBAR = "#111827"
COR_CARD = "#161b22"
COR_AZUL_PRIMARIO = "#1f6feb"
COR_AZUL_HOVER = "#388bfd"
COR_AZUL_ESCURO = "#0d419d"
COR_TEXTO = "#e6edf3"
COR_TEXTO_SECUNDARIO = "#8b949e"
COR_BORDA = "#21262d"
COR_VERMELHO = "#8b1e1e"
COR_VERMELHO_HOVER = "#b02a2a"

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")


def estilizar_treeview():
    estilo = ttk.Style()
    estilo.theme_use("default")
    estilo.configure("Treeview", background=COR_FUNDO, foreground=COR_TEXTO,
                      fieldbackground=COR_FUNDO, rowheight=30, borderwidth=0,
                      font=("Segoe UI", 11))
    estilo.configure("Treeview.Heading", background=COR_AZUL_ESCURO, foreground=COR_TEXTO,
                      font=("Segoe UI", 11, "bold"), borderwidth=0)
    estilo.map("Treeview", background=[("selected", COR_AZUL_PRIMARIO)])


#Login e Cadastro de Usuario
class TelaLogin(ctk.CTkToplevel if False else ctk.CTk):
    def __init__(self, usuarios: Usuarios, ao_logar):
        super().__init__()
        self.usuarios = usuarios
        self.ao_logar = ao_logar

        self.title("Login - Sistema de Estoque")
        self.geometry("420x480")
        self.configure(fg_color=COR_FUNDO)
        self.resizable(False, False)

        card = ctk.CTkFrame(self, fg_color=COR_CARD, corner_radius=14,
                             border_width=1, border_color=COR_BORDA)
        card.pack(expand=True, fill="both", padx=25, pady=25)

        ctk.CTkLabel(card, text="📦 Sistema de Estoque", font=ctk.CTkFont(size=20, weight="bold"),
                     text_color=COR_TEXTO).pack(pady=(30, 5))
        ctk.CTkLabel(card, text="Entre com sua conta", font=ctk.CTkFont(size=13),
                     text_color=COR_TEXTO_SECUNDARIO).pack(pady=(0, 20))

        self.entry_user = self._campo(card, "Usuário")
        self.entry_senha = self._campo(card, "Senha", mostrar="*")

        btn_entrar = ctk.CTkButton(card, text="Entrar", command=self.fazer_login,
                                    height=42, corner_radius=8, fg_color=COR_AZUL_PRIMARIO,
                                    hover_color=COR_AZUL_HOVER, font=ctk.CTkFont(size=14, weight="bold"))
        btn_entrar.pack(pady=(20, 10), padx=30, fill="x")

        btn_criar = ctk.CTkButton(card, text="Criar nova conta", command=self.abrir_cadastro,
                                   height=38, corner_radius=8, fg_color="transparent",
                                   hover_color=COR_BORDA, text_color=COR_AZUL_PRIMARIO,
                                   border_width=1, border_color=COR_AZUL_PRIMARIO)
        btn_criar.pack(padx=30, fill="x")

        self.bind("<Return>", lambda e: self.fazer_login())

    def _campo(self, master, label, mostrar=None):
        ctk.CTkLabel(master, text=label, text_color=COR_TEXTO_SECUNDARIO,
                     font=ctk.CTkFont(size=12)).pack(anchor="w", padx=30, pady=(8, 2))
        entry = ctk.CTkEntry(master, height=38, corner_radius=8, fg_color=COR_FUNDO,
                              border_color=COR_AZUL_PRIMARIO, text_color=COR_TEXTO,
                              show=mostrar)
        entry.pack(padx=30, fill="x")
        return entry

    def fazer_login(self):
        username = self.entry_user.get().strip()
        senha = self.entry_senha.get().strip()
        if not username or not senha:
            messagebox.showwarning("Campos obrigatórios", "Preencha usuário e senha.")
            return
        usuario = self.usuarios.login(username, senha)
        if usuario:
            self.destroy()
            self.ao_logar(usuario)
        else:
            messagebox.showerror("Erro", "Usuário ou senha incorretos.")

    def abrir_cadastro(self):
        JanelaCadastroUsuario(self, self.usuarios)


class JanelaCadastroUsuario(ctk.CTkToplevel):
    def __init__(self, master, usuarios: Usuarios):
        super().__init__(master)
        self.usuarios = usuarios
        self.title("Criar Conta")
        self.geometry("380x420")
        self.configure(fg_color=COR_FUNDO)
        self.resizable(False, False)
        self.grab_set()

        ctk.CTkLabel(self, text="Criar Nova Conta", font=ctk.CTkFont(size=18, weight="bold"),
                     text_color=COR_TEXTO).pack(pady=(25, 15))

        self.entry_user = self._campo("Usuário")
        self.entry_senha = self._campo("Senha", mostrar="*")
        self.entry_email = self._campo("E-mail (alertas de estoque)")

        ctk.CTkButton(self, text="Cadastrar", command=self.cadastrar,
                      height=40, corner_radius=8, fg_color=COR_AZUL_PRIMARIO,
                      hover_color=COR_AZUL_HOVER).pack(pady=25, padx=30, fill="x")

    def _campo(self, label, mostrar=None):
        ctk.CTkLabel(self, text=label, text_color=COR_TEXTO_SECUNDARIO,
                     font=ctk.CTkFont(size=12)).pack(anchor="w", padx=30, pady=(8, 2))
        entry = ctk.CTkEntry(self, height=38, corner_radius=8, fg_color=COR_CARD,
                              border_color=COR_AZUL_PRIMARIO, text_color=COR_TEXTO, show=mostrar)
        entry.pack(padx=30, fill="x")
        return entry

    def cadastrar(self):
        username = self.entry_user.get().strip()
        senha = self.entry_senha.get().strip()
        email = self.entry_email.get().strip()
        if not username or not senha or not email:
            messagebox.showwarning("Campos obrigatórios", "Preencha todos os campos.")
            return
        if self.usuarios.cadastrar_usuario(username, senha, email):
            messagebox.showinfo("Sucesso", "Conta criada! Agora você pode entrar.")
            self.destroy()
        else:
            messagebox.showerror("Erro", "Esse nome de usuário já existe.")


#Principal
class App(ctk.CTk):
    def __init__(self, gestao: Gestao, vendas: Vendas, usuario_logado):
        super().__init__()
        self.gestao = gestao
        self.vendas = vendas
        self.usuario = usuario_logado

        self.title("Sistema de Gestão de Estoque")
        self.geometry("1100x650")
        self.configure(fg_color=COR_FUNDO)
        self.minsize(950, 600)

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        estilizar_treeview()

        self._criar_sidebar()
        self._criar_area_principal()

        self.frame_estoque.tkraise()
        self.atualizar_tabela_estoque()

    #Barra lateral
    def _criar_sidebar(self):
        sidebar = ctk.CTkFrame(self, width=230, corner_radius=0, fg_color=COR_SIDEBAR)
        sidebar.grid(row=0, column=0, sticky="nswe")
        sidebar.grid_propagate(False)

        ctk.CTkLabel(sidebar, text="📦 Gestão de\nEstoque", font=ctk.CTkFont(size=20, weight="bold"),
                     text_color=COR_TEXTO, justify="left").pack(pady=(30, 5), padx=20, anchor="w")
        ctk.CTkLabel(sidebar, text=f"Olá, {self.usuario['username']}", font=ctk.CTkFont(size=12),
                     text_color=COR_TEXTO_SECUNDARIO).pack(pady=(0, 30), padx=20, anchor="w")

        botoes = [
            ("📦  Estoque", self.mostrar_estoque),
            ("💰  Vendas", self.mostrar_vendas),
            ("🚚  Fornecedores", self.mostrar_fornecedores),
            ("📊  Relatório", self.mostrar_relatorio),
        ]
        for texto, comando in botoes:
            ctk.CTkButton(sidebar, text=texto, command=comando, anchor="w", height=42,
                          corner_radius=8, fg_color="transparent", hover_color=COR_AZUL_ESCURO,
                          text_color=COR_TEXTO, font=ctk.CTkFont(size=14)).pack(pady=4, padx=15, fill="x")

        ctk.CTkLabel(sidebar, text="v1.0 • CustomTkinter", text_color=COR_TEXTO_SECUNDARIO,
                     font=ctk.CTkFont(size=11)).pack(side="bottom", pady=20)

    #Area principal
    def _criar_area_principal(self):
        container = ctk.CTkFrame(self, fg_color=COR_FUNDO)
        container.grid(row=0, column=1, sticky="nswe", padx=20, pady=20)
        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        self.frame_estoque = self._criar_frame_estoque(container)
        self.frame_vendas = self._criar_frame_vendas(container)
        self.frame_fornecedores = self._criar_frame_fornecedores(container)
        self.frame_relatorio = self._criar_frame_relatorio(container)

        for frame in (self.frame_estoque, self.frame_vendas, self.frame_fornecedores, self.frame_relatorio):
            frame.grid(row=0, column=0, sticky="nswe")

    def _card_header(self, frame, titulo):
        header = ctk.CTkFrame(frame, fg_color="transparent")
        header.pack(fill="x", padx=25, pady=(20, 10))
        ctk.CTkLabel(header, text=titulo, font=ctk.CTkFont(size=20, weight="bold"),
                     text_color=COR_TEXTO).pack(side="left")
        return header

    #Estoque
    def _criar_frame_estoque(self, master):
        frame = ctk.CTkFrame(master, fg_color=COR_CARD, corner_radius=12,
                              border_width=1, border_color=COR_BORDA)
        self._card_header(frame, "Estoque")

        form = ctk.CTkFrame(frame, fg_color="transparent")
        form.pack(fill="x", padx=25, pady=(0, 10))
        form.grid_columnconfigure((0, 1, 2), weight=1)

        self.entry_produto = self._campo_grid(form, "Produto", 0, 0)
        self.entry_qtd = self._campo_grid(form, "Quantidade", 0, 1)
        self.entry_preco = self._campo_grid(form, "Preço (ex: 10.50)", 0, 2)

        btns = ctk.CTkFrame(frame, fg_color="transparent")
        btns.pack(fill="x", padx=25, pady=(0, 15))
        ctk.CTkButton(btns, text="➕ Adicionar / Atualizar", command=self.adicionar_produto,
                      fg_color=COR_AZUL_PRIMARIO, hover_color=COR_AZUL_HOVER, height=38).pack(side="left")
        ctk.CTkButton(btns, text="🗑 Remover selecionado", command=self.remover_produto_selecionado,
                      fg_color=COR_VERMELHO, hover_color=COR_VERMELHO_HOVER, height=38).pack(side="left", padx=10)
        ctk.CTkButton(btns, text="🔄 Atualizar lista", command=self.atualizar_tabela_estoque,
                      fg_color=COR_AZUL_ESCURO, hover_color=COR_AZUL_PRIMARIO, height=38).pack(side="left")

        tabela_frame = ctk.CTkFrame(frame, fg_color="transparent")
        tabela_frame.pack(fill="both", expand=True, padx=25, pady=(0, 20))

        colunas = ("ID", "Produto", "Quantidade", "Preço")
        self.tabela_estoque = ttk.Treeview(tabela_frame, columns=colunas, show="headings")
        for col in colunas:
            self.tabela_estoque.heading(col, text=col)
            self.tabela_estoque.column(col, anchor="w", width=120)
        self.tabela_estoque.pack(fill="both", expand=True)

        return frame

    def _campo_grid(self, master, label, row, col):
        wrap = ctk.CTkFrame(master, fg_color="transparent")
        wrap.grid(row=row, column=col, sticky="we", padx=6)
        ctk.CTkLabel(wrap, text=label, text_color=COR_TEXTO_SECUNDARIO,
                     font=ctk.CTkFont(size=12)).pack(anchor="w")
        entry = ctk.CTkEntry(wrap, height=36, corner_radius=8, fg_color=COR_FUNDO,
                              border_color=COR_AZUL_PRIMARIO, text_color=COR_TEXTO)
        entry.pack(fill="x")
        return entry

    def adicionar_produto(self):
        produto = self.entry_produto.get().strip()
        qtd_txt = self.entry_qtd.get().strip()
        preco_txt = self.entry_preco.get().strip()

        if not produto or not qtd_txt:
            messagebox.showwarning("Campos obrigatórios", "Informe produto e quantidade.")
            return
        try:
            quantidade = int(qtd_txt)
            preco = float(preco_txt) if preco_txt else None
        except ValueError:
            messagebox.showerror("Erro", "Quantidade deve ser inteira e preço numérico (ex: 10.50).")
            return

        nova_qtd = self.gestao.adicionar_produto(produto, quantidade, preco)
        messagebox.showinfo("Sucesso", f"'{produto}' atualizado. Nova quantidade: {nova_qtd}")

        verificar_e_alertar(self.gestao, produto, self.usuario["email"])

        for e in (self.entry_produto, self.entry_qtd, self.entry_preco):
            e.delete(0, "end")
        self.atualizar_tabela_estoque()

    def remover_produto_selecionado(self):
        sel = self.tabela_estoque.selection()
        if not sel:
            messagebox.showwarning("Nada selecionado", "Selecione um produto na tabela.")
            return
        valores = self.tabela_estoque.item(sel[0], "values")
        produto = valores[1]
        if messagebox.askyesno("Confirmar", f"Remover '{produto}' do estoque?"):
            self.gestao.remover_produto(produto)
            self.atualizar_tabela_estoque()

    def atualizar_tabela_estoque(self):
        for item in self.tabela_estoque.get_children():
            self.tabela_estoque.delete(item)
        for id_p, produto, qtd, preco in self.gestao.lista_estoque_completa():
            self.tabela_estoque.insert("", "end", values=(id_p, produto, qtd, f"R$ {preco:.2f}"))
        self._atualizar_combos_produtos()

    #Vendas
    def _criar_frame_vendas(self, master):
        frame = ctk.CTkFrame(master, fg_color=COR_CARD, corner_radius=12,
                              border_width=1, border_color=COR_BORDA)
        self._card_header(frame, "Vendas")

        form = ctk.CTkFrame(frame, fg_color="transparent")
        form.pack(fill="x", padx=25, pady=(0, 10))
        form.grid_columnconfigure((0, 1), weight=1)

        wrap1 = ctk.CTkFrame(form, fg_color="transparent")
        wrap1.grid(row=0, column=0, sticky="we", padx=6)
        ctk.CTkLabel(wrap1, text="Produto", text_color=COR_TEXTO_SECUNDARIO,
                     font=ctk.CTkFont(size=12)).pack(anchor="w")
        self.combo_produto_venda = ctk.CTkComboBox(wrap1, values=[], height=36, fg_color=COR_FUNDO,
                                                     border_color=COR_AZUL_PRIMARIO, text_color=COR_TEXTO,
                                                     button_color=COR_AZUL_PRIMARIO,
                                                     button_hover_color=COR_AZUL_HOVER)
        self.combo_produto_venda.pack(fill="x")

        self.entry_qtd_venda = self._campo_grid(form, "Quantidade", 0, 1)

        btns = ctk.CTkFrame(frame, fg_color="transparent")
        btns.pack(fill="x", padx=25, pady=(10, 15))
        ctk.CTkButton(btns, text="💰 Registrar Venda", command=self.registrar_venda,
                      fg_color=COR_AZUL_PRIMARIO, hover_color=COR_AZUL_HOVER, height=38).pack(side="left")
        ctk.CTkButton(btns, text="🗑 Remover venda selecionada", command=self.remover_venda_selecionada,
                      fg_color=COR_VERMELHO, hover_color=COR_VERMELHO_HOVER, height=38).pack(side="left", padx=10)
        ctk.CTkButton(btns, text="🔄 Atualizar lista", command=self.atualizar_tabela_vendas,
                      fg_color=COR_AZUL_ESCURO, hover_color=COR_AZUL_PRIMARIO, height=38).pack(side="left")

        tabela_frame = ctk.CTkFrame(frame, fg_color="transparent")
        tabela_frame.pack(fill="both", expand=True, padx=25, pady=(0, 20))

        colunas = ("ID", "Produto", "Qtd", "Preço Unit.", "Total", "Data")
        self.tabela_vendas = ttk.Treeview(tabela_frame, columns=colunas, show="headings")
        for col in colunas:
            self.tabela_vendas.heading(col, text=col)
            self.tabela_vendas.column(col, anchor="w", width=110)
        self.tabela_vendas.pack(fill="both", expand=True)

        return frame

    def registrar_venda(self):
        produto = self.combo_produto_venda.get().strip()
        qtd_txt = self.entry_qtd_venda.get().strip()

        if not produto or not qtd_txt:
            messagebox.showwarning("Campos obrigatórios", "Selecione o produto e informe a quantidade.")
            return
        try:
            quantidade = int(qtd_txt)
        except ValueError:
            messagebox.showerror("Erro", "Quantidade deve ser um número inteiro.")
            return

        sucesso, resultado = self.vendas.registrar_venda(produto, quantidade)
        if sucesso:
            messagebox.showinfo("Venda registrada", f"Total: R$ {resultado:.2f}")
            verificar_e_alertar(self.gestao, produto, self.usuario["email"])
            self.entry_qtd_venda.delete(0, "end")
            self.atualizar_tabela_vendas()
            self.atualizar_tabela_estoque()
        else:
            messagebox.showerror("Erro na venda", resultado)

    def remover_venda_selecionada(self):
        sel = self.tabela_vendas.selection()
        if not sel:
            messagebox.showwarning("Nada selecionado", "Selecione uma venda na tabela.")
            return
        venda_id = self.tabela_vendas.item(sel[0], "values")[0]
        if messagebox.askyesno("Confirmar", "Remover esta venda? O produto voltará ao estoque."):
            self.vendas.remover_venda(int(venda_id))
            self.atualizar_tabela_vendas()
            self.atualizar_tabela_estoque()

    def atualizar_tabela_vendas(self):
        for item in self.tabela_vendas.get_children():
            self.tabela_vendas.delete(item)
        for id_v, produto, qtd, preco_unit, total, data in self.vendas.listar_vendas():
            self.tabela_vendas.insert("", "end", values=(id_v, produto, qtd, f"R$ {preco_unit:.2f}",
                                                          f"R$ {total:.2f}", data))

    def _atualizar_combos_produtos(self):
        produtos = [p[0] for p in self.gestao.lista_estoque()]
        if hasattr(self, "combo_produto_venda"):
            self.combo_produto_venda.configure(values=produtos)

    #Fornecedores
    def _criar_frame_fornecedores(self, master):
        frame = ctk.CTkFrame(master, fg_color=COR_CARD, corner_radius=12,
                              border_width=1, border_color=COR_BORDA)
        self._card_header(frame, "Fornecedores")

        form = ctk.CTkFrame(frame, fg_color="transparent")
        form.pack(fill="x", padx=25, pady=(0, 10))
        form.grid_columnconfigure((0, 1, 2), weight=1)

        self.entry_forn_nome = self._campo_grid(form, "Nome", 0, 0)
        self.entry_forn_celular = self._campo_grid(form, "Celular", 0, 1)
        self.entry_forn_produto = self._campo_grid(form, "Produto fornecido", 0, 2)

        btns = ctk.CTkFrame(frame, fg_color="transparent")
        btns.pack(fill="x", padx=25, pady=(10, 15))
        ctk.CTkButton(btns, text="➕ Cadastrar fornecedor", command=self.cadastrar_fornecedor,
                      fg_color=COR_AZUL_PRIMARIO, hover_color=COR_AZUL_HOVER, height=38).pack(side="left")
        ctk.CTkButton(btns, text="🗑 Remover selecionado", command=self.remover_fornecedor_selecionado,
                      fg_color=COR_VERMELHO, hover_color=COR_VERMELHO_HOVER, height=38).pack(side="left", padx=10)
        ctk.CTkButton(btns, text="🔄 Atualizar lista", command=self.atualizar_tabela_fornecedores,
                      fg_color=COR_AZUL_ESCURO, hover_color=COR_AZUL_PRIMARIO, height=38).pack(side="left")

        tabela_frame = ctk.CTkFrame(frame, fg_color="transparent")
        tabela_frame.pack(fill="both", expand=True, padx=25, pady=(0, 20))

        colunas = ("ID", "Nome", "Celular", "Produto")
        self.tabela_fornecedores = ttk.Treeview(tabela_frame, columns=colunas, show="headings")
        for col in colunas:
            self.tabela_fornecedores.heading(col, text=col)
            self.tabela_fornecedores.column(col, anchor="w", width=130)
        self.tabela_fornecedores.pack(fill="both", expand=True)

        return frame

    def cadastrar_fornecedor(self):
        nome = self.entry_forn_nome.get().strip()
        celular = self.entry_forn_celular.get().strip()
        produto = self.entry_forn_produto.get().strip()
        if not nome or not produto:
            messagebox.showwarning("Campos obrigatórios", "Informe ao menos nome e produto.")
            return
        self.gestao.adicionar_fornecedor(nome, celular, produto)
        for e in (self.entry_forn_nome, self.entry_forn_celular, self.entry_forn_produto):
            e.delete(0, "end")
        self.atualizar_tabela_fornecedores()

    def remover_fornecedor_selecionado(self):
        sel = self.tabela_fornecedores.selection()
        if not sel:
            messagebox.showwarning("Nada selecionado", "Selecione um fornecedor na tabela.")
            return
        forn_id = self.tabela_fornecedores.item(sel[0], "values")[0]
        if messagebox.askyesno("Confirmar", "Remover este fornecedor?"):
            self.gestao.remover_fornecedor(int(forn_id))
            self.atualizar_tabela_fornecedores()

    def atualizar_tabela_fornecedores(self):
        for item in self.tabela_fornecedores.get_children():
            self.tabela_fornecedores.delete(item)
        for id_f, nome, celular, produto in self.gestao.listar_fornecedores():
            self.tabela_fornecedores.insert("", "end", values=(id_f, nome, celular, produto))

    #Relatorio
    def _criar_frame_relatorio(self, master):
        frame = ctk.CTkFrame(master, fg_color=COR_CARD, corner_radius=12,
                              border_width=1, border_color=COR_BORDA)
        self._card_header(frame, "Relatório")

        self.label_valor_total = ctk.CTkLabel(frame, text="Valor total do estoque: R$ 0,00",
                                               font=ctk.CTkFont(size=16, weight="bold"),
                                               text_color=COR_TEXTO)
        self.label_valor_total.pack(pady=20, padx=25, anchor="w")

        ctk.CTkButton(frame, text="📊 Gerar relatório em Excel (com gráfico)",
                      command=self.gerar_relatorio, height=42, corner_radius=8,
                      fg_color=COR_AZUL_PRIMARIO, hover_color=COR_AZUL_HOVER).pack(padx=25, pady=10, anchor="w")

        ctk.CTkButton(frame, text="🔄 Atualizar valor total", command=self.atualizar_valor_total,
                      height=38, corner_radius=8, fg_color=COR_AZUL_ESCURO,
                      hover_color=COR_AZUL_PRIMARIO).pack(padx=25, pady=5, anchor="w")

        return frame

    def atualizar_valor_total(self):
        total = self.gestao.valor_total_estoque()
        self.label_valor_total.configure(text=f"Valor total do estoque: R$ {total:.2f}")

    def gerar_relatorio(self):
        sucesso = self.gestao.gerar_grafico_estoque("relatorio_estoque.xlsx")
        if sucesso:
            messagebox.showinfo("Relatório gerado", "Arquivo 'relatorio_estoque.xlsx' criado com sucesso.")
        else:
            messagebox.showwarning("Estoque vazio", "Não há produtos em estoque para gerar o relatório.")
        self.atualizar_valor_total()

    #Navegacao
    def mostrar_estoque(self):
        self.atualizar_tabela_estoque()
        self.frame_estoque.tkraise()

    def mostrar_vendas(self):
        self._atualizar_combos_produtos()
        self.atualizar_tabela_vendas()
        self.frame_vendas.tkraise()

    def mostrar_fornecedores(self):
        self.atualizar_tabela_fornecedores()
        self.frame_fornecedores.tkraise()

    def mostrar_relatorio(self):
        self.atualizar_valor_total()
        self.frame_relatorio.tkraise()


#Inicializacao
def iniciar_app_principal(usuario_logado):
    gestao = Gestao("estoque.db")
    vendas = Vendas("estoque.db", gestao)
    app = App(gestao, vendas, usuario_logado)
    app.mainloop()


if __name__ == "__main__":
    usuarios = Usuarios("estoque.db")
    tela_login = TelaLogin(usuarios, ao_logar=iniciar_app_principal)
    tela_login.mainloop()
